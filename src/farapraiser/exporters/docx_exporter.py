"""DOCX document exporter."""

from typing import List
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from farapraiser.parsers.base import DocumentBlock, BlockType
from farapraiser.exporters.base import DocumentExporter


class DOCXExporter(DocumentExporter):
    """Exporter for Microsoft Word DOCX files."""
    
    def export(self, blocks: List[DocumentBlock], output_path: str) -> None:
        """Export blocks to DOCX file."""
        document = Document()
        
        for block in blocks:
            if block.block_type == BlockType.HEADING:
                para = document.add_heading(block.content, level=self._get_heading_level(block))
            
            elif block.block_type == BlockType.PARAGRAPH:
                para = document.add_paragraph(block.content)
                para.style = block.metadata.get("style", "Normal")
            
            elif block.block_type == BlockType.LIST:
                for item in block.content:
                    para = document.add_paragraph(item, style='List Bullet')
            
            elif block.block_type == BlockType.TABLE:
                if block.content:
                    num_rows = len(block.content)
                    num_cols = len(block.content[0]) if num_rows > 0 else 0
                    
                    if num_rows > 0 and num_cols > 0:
                        table = document.add_table(rows=num_rows, cols=num_cols)
                        table.style = 'Light Grid Accent 1'
                        
                        for i, row_data in enumerate(block.content):
                            for j, cell_data in enumerate(row_data):
                                table.rows[i].cells[j].text = str(cell_data)
        
        document.save(output_path)
    
    def _get_heading_level(self, block: DocumentBlock) -> int:
        """Extract heading level from metadata."""
        style = block.metadata.get("style", "") or block.metadata.get("level", "")
        if "Heading 1" in style or "heading1" in style.lower():
            return 1
        elif "Heading 2" in style or "heading2" in style.lower():
            return 2
        elif "Heading 3" in style or "heading3" in style.lower():
            return 3
        return 1
