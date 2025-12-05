"""DOCX document parser."""

from typing import List
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from farapraiser.parsers.base import DocumentParser, DocumentBlock, BlockType


class DOCXParser(DocumentParser):
    """Parser for Microsoft Word DOCX files."""
    
    def supports_format(self, file_path: str) -> bool:
        """Check if file is DOCX format."""
        return file_path.lower().endswith('.docx')
    
    def parse(self, file_path: str) -> List[DocumentBlock]:
        """Parse DOCX file and extract blocks."""
        document = Document(file_path)
        blocks = []
        index = 0
        
        for element in document.element.body:
            # Handle paragraphs
            if element.tag.endswith('p'):
                para = None
                for p in document.paragraphs:
                    if p._element == element:
                        para = p
                        break
                
                if para and para.text.strip():
                    style_name = para.style.name if para.style else "Normal"
                    
                    # Detect headings
                    if 'Heading' in style_name:
                        block_type = BlockType.HEADING
                        metadata = {"level": style_name, "style": style_name}
                    else:
                        block_type = BlockType.PARAGRAPH
                        metadata = {"style": style_name}
                    
                    blocks.append(DocumentBlock(
                        block_type=block_type,
                        content=para.text,
                        metadata=metadata,
                        index=index
                    ))
                    index += 1
            
            # Handle tables
            elif element.tag.endswith('tbl'):
                table = None
                for t in document.tables:
                    if t._element == element:
                        table = t
                        break
                
                if table:
                    table_data = []
                    for row in table.rows:
                        row_data = [cell.text for cell in row.cells]
                        table_data.append(row_data)
                    
                    blocks.append(DocumentBlock(
                        block_type=BlockType.TABLE,
                        content=table_data,
                        metadata={"rows": len(table.rows), "cols": len(table.columns)},
                        index=index
                    ))
                    index += 1
        
        return blocks
